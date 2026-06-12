import io
import json
import base64
import logging
import textwrap
from django.shortcuts import render, get_object_or_404
from django.views.generic import TemplateView, View
from django.contrib.auth.mixins import LoginRequiredMixin
from django.http import JsonResponse, HttpResponse
from django.utils.decorators import method_decorator
from django.views.decorators.csrf import csrf_exempt
from django.conf import settings
from django.db import transaction
from django.utils import timezone
from openai import OpenAI

# ReportLab imports for PDF generation
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import mm, cm
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT, TA_JUSTIFY
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
    HRFlowable, Image as RLImage, KeepTogether, PageBreak
)
from reportlab.graphics.shapes import Drawing
from reportlab.pdfgen import canvas
from PIL import Image as PILImage

from apps.meetings.models import Meeting, MeetingMinute
from apps.communities.models import Community, Membership
from apps.global_data.enum import MinuteStatus
from apps.finance.utils.admin_mixins import AdminSeasonMixin
from django.utils.decorators import method_decorator
from apps.log.decorators import log_activity_and_errors

logger = logging.getLogger(__name__)

class AdminMeetingMinuteView(AdminSeasonMixin, LoginRequiredMixin, TemplateView):
    template_name = 'publics/admin/meetings_minutes/meetings_minutes.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        community_id = self.kwargs.get('community_id')
        community = Community.objects.get(id=community_id)
        context['community'] = community
        
        # Fetch last 3 minutes for initial display
        recent_minutes = MeetingMinute.objects.filter(
            meeting__community=community
        ).select_related('meeting', 'prepared_by').order_by('-created')[:3]
        
        context['recent_minutes'] = recent_minutes
        
        # Total count for pagination display (e.g., "3 rows", or for "Load More" logic)
        context['total_minutes_count'] = MeetingMinute.objects.filter(
            meeting__community=community
        ).count()
        
        return context

class MeetingMinuteAIView(LoginRequiredMixin, TemplateView):
    """Handles AI logic: Transcription and Generation"""
    
    @method_decorator(log_activity_and_errors())
    def post(self, request, *args, **kwargs):
        action = request.POST.get('action')
        client = OpenAI(api_key=settings.OPENAI_API_KEY)

        if action == 'transcribe':
            audio_file = request.FILES.get('audio')
            if not audio_file:
                return JsonResponse({'status': 'error', 'message': 'No file provided'})
            
            ext = audio_file.name.split('.')[-1].lower()
            audio_exts = ['flac', 'm4a', 'mp3', 'mp4', 'mpeg', 'mpga', 'oga', 'ogg', 'wav', 'webm']
            
            try:
                if ext in audio_exts:
                    # Use OpenAI Whisper for audio
                    transcription = client.audio.transcriptions.create(
                        model="whisper-1", 
                        file=(audio_file.name, audio_file.read())
                    )
                    return JsonResponse({'status': 'success', 'transcript': transcription.text})
                
                elif ext == 'pdf':
                    import pypdf
                    reader = pypdf.PdfReader(audio_file)
                    text = ""
                    for page in reader.pages:
                        text += page.extract_text() + "\n"
                    return JsonResponse({'status': 'success', 'transcript': text.strip()})
                
                elif ext == 'docx':
                    import docx
                    doc = docx.Document(audio_file)
                    text = "\n".join([para.text for para in doc.paragraphs])
                    return JsonResponse({'status': 'success', 'transcript': text.strip()})
                
                elif ext == 'txt':
                    text = audio_file.read().decode('utf-8')
                    return JsonResponse({'status': 'success', 'transcript': text.strip()})
                
                else:
                    return JsonResponse({'status': 'error', 'message': f'Unsupported file format: {ext}'})
                    
            except Exception as e:
                logger.error(f"Transcription/Extraction error: {e}")
                return JsonResponse({'status': 'error', 'message': f"Failed to process {ext} file: {str(e)}"})

        elif action == 'generate':
            transcript = request.POST.get('transcript')
            language = request.POST.get('language', 'English')
            
            if not transcript:
                return JsonResponse({'status': 'error', 'message': 'No transcript provided'})

            try:
                # Use GPT-4 to generate summary and professional minutes
                prompt = f"""
                You are a professional secretary. Based on the following meeting transcript in {language}, 
                please generate exactly two things:
                1. A concise meeting summary.
                2. Professional meeting minutes with sections for attendees (if mentioned), agenda, discussions, and decisions.

                Format your response as a JSON object with keys 'summary' and 'minutes'. 
                IMPORTANT: Both values must be strings. The 'minutes' should be well-formatted with newlines and bullet points.
                
                Transcript:
                {transcript}
                """
                
                response = client.chat.completions.create(
                    model="gpt-4o",
                    messages=[{"role": "user", "content": prompt}],
                    response_format={"type": "json_object"}
                )
                
                ai_content = json.loads(response.choices[0].message.content)
                return JsonResponse({
                    'status': 'success', 
                    'summary': ai_content.get('summary'), 
                    'minutes': ai_content.get('minutes')
                })
            except Exception as e:
                logger.error(f"AI Generation error: {e}")
                return JsonResponse({'status': 'error', 'message': str(e)})

        return JsonResponse({'status': 'error', 'message': 'Invalid action'})

class MeetingAttendanceAPIView(LoginRequiredMixin, View):
    """Handles saving/updating and fetching attendance for a meeting"""
    
    def get(self, request, *args, **kwargs):
        try:
            meeting_id = request.GET.get('meeting_id')
            community_id = request.GET.get('community_id')
            date = request.GET.get('date')
            title = request.GET.get('title')

            meeting = None
            if meeting_id:
                meeting = get_object_or_404(Meeting, id=meeting_id)
            elif all([community_id, date, title]):
                meeting = Meeting.objects.filter(
                    community_id=community_id,
                    scheduled_date=date,
                    title=title
                ).first()

            if not meeting:
                # If meeting doesn't exist, just return exists: False
                # The frontend will then fetch members via the normal membership API
                return JsonResponse({'status': 'success', 'exists': False})

            from apps.communities.models import Membership
            from apps.meetings.models import MeetingAttendance

            # Get all active members
            memberships = Membership.objects.filter(community=meeting.community, status='active').select_related('user')
            # Get existing attendance
            attendance_map = {str(a.membership_id): a.was_present for a in MeetingAttendance.objects.filter(meeting=meeting)}

            data = []
            for m in memberships:
                data.append({
                    'id': str(m.id),
                    'name': m.user.get_full_name or m.user.username,
                    'was_present': attendance_map.get(str(m.id), True) # Default to True
                })

            return JsonResponse({
                'status': 'success',
                'exists': True,
                'meeting_id': str(meeting.id),
                'attendance': data
            })
        except Exception as e:
            logger.error(f"Attendance fetch error: {e}")
            return JsonResponse({'status': 'error', 'message': str(e)}, status=400)

    @method_decorator(log_activity_and_errors())
    def post(self, request, *args, **kwargs):
        try:
            data = json.loads(request.body)
            community_id = data.get('community_id')
            date = data.get('date')
            title = data.get('title')
            attendance_list = data.get('attendance', [])

            if not all([community_id, date, title]):
                return JsonResponse({'status': 'error', 'message': 'Missing required fields (Community, Date, Title)'}, status=400)

            community = get_object_or_404(Community, id=community_id)

            from apps.meetings.models import MeetingAttendance
            from apps.communities.models import Membership

            with transaction.atomic():
                # 1. Find or create meeting
                meeting, created = Meeting.objects.get_or_create(
                    community=community,
                    scheduled_date=date,
                    title=title,
                    defaults={
                        'start_time': timezone.now().time(),
                        'end_time': timezone.now().time(),
                        'chaired_by': request.user
                    }
                )

                # 2. Save attendance records
                for record in attendance_list:
                    membership_id = record.get('membership_id')
                    was_present = record.get('was_present', True)
                    
                    membership = get_object_or_404(Membership, id=membership_id)
                    MeetingAttendance.objects.update_or_create(
                        meeting=meeting,
                        membership=membership,
                        defaults={'was_present': was_present}
                    )

            return JsonResponse({
                'status': 'success', 
                'message': 'Attendance saved successfully!', 
                'meeting_id': str(meeting.id)
            })
        except Exception as e:
            logger.error(f"Attendance save error: {e}")
            return JsonResponse({'status': 'error', 'message': str(e)}, status=400)

class MeetingMinuteSaveView(LoginRequiredMixin, TemplateView):
    """Handles saving meeting and minutes to database"""

    @method_decorator(log_activity_and_errors())
    def post(self, request, *args, **kwargs):
        try:
            data = request.POST
            title = data.get('title')
            date = data.get('date')
            language = data.get('language')
            transcript = data.get('transcript')
            summary = data.get('summary')
            minutes_content = data.get('minutes')
            community_id = data.get('community_id')
            meeting_id = data.get('meeting_id')

            if not all([title, date, community_id]):
                return JsonResponse({'status': 'error', 'message': 'Missing required fields (Title, Date, Community)'})

            community = Community.objects.get(id=community_id)

            with transaction.atomic():
                # 1. Create/Update Meeting
                if meeting_id:
                    meeting = get_object_or_404(Meeting, id=meeting_id)
                else:
                    meeting = Meeting.objects.create(
                        community=community,
                        title=title,
                        scheduled_date=date,
                        start_time=timezone.now().time(),
                        end_time=timezone.now().time(), # Placeholder
                        chaired_by=request.user
                    )

                # 2. Create/Update MeetingMinute
                MeetingMinute.objects.update_or_create(
                    meeting=meeting,
                    defaults={
                        'title': title,
                        'minute_date': date,
                        'start_time': meeting.start_time,
                        'end_time': meeting.end_time,
                        'prepared_by': request.user,
                        'transcript': transcript,
                        'summary': summary,
                        'discussions': minutes_content,
                        'language': language,
                        'status': MinuteStatus.DRAFT
                    }
                )

            return JsonResponse({'status': 'success', 'message': 'Meeting minutes saved successfully!'})
        except Exception as e:
            logger.error(f"Save error: {e}")
            return JsonResponse({'status': 'error', 'message': str(e)})

class MeetingMinuteDetailView(LoginRequiredMixin, TemplateView):
    """Fetches details of a specific meeting minute"""
    def get(self, request, pk, *args, **kwargs):
        try:
            minute = MeetingMinute.objects.get(pk=pk)
            return JsonResponse({
                'status': 'success',
                'title': minute.title,
                'date': minute.minute_date.isoformat() if minute.minute_date else '',
                'language': minute.language,
                'transcript': minute.transcript,
                'summary': minute.summary,
                'minutes': minute.discussions,
                'meeting_id': str(minute.meeting_id)
            })
        except MeetingMinute.DoesNotExist:
            return JsonResponse({'status': 'error', 'message': 'Minute not found'})
        except Exception as e:
            return JsonResponse({'status': 'error', 'message': str(e)})

class AllMeetingMinutesAPIView(LoginRequiredMixin, View):
    """API to fetch all meeting minutes for a community"""
    def get(self, request, community_id, *args, **kwargs):
        try:
            community = get_object_or_404(Community, id=community_id)
            minutes = MeetingMinute.objects.filter(
                meeting__community=community
            ).select_related('meeting', 'prepared_by').order_by('-created')
            
            data = []
            for m in minutes:
                data.append({
                    'id': str(m.id),
                    'minute_date': m.minute_date.isoformat() if m.minute_date else '---',
                    'title': m.title or 'Meeting Minutes',
                    'audio_file': bool(m.audio_file),
                    'prepared_by': m.prepared_by.get_full_name or m.prepared_by.email if m.prepared_by else '---',
                    'created_at': m.created.strftime('%Y-%m-%d %H:%M:%S'),
                    'pdf_url': f'/meetings/meeting-minutes/pdf/{m.id}/'
                })
                
            return JsonResponse({'status': 'success', 'minutes': data})
        except Exception as e:
            return JsonResponse({'status': 'error', 'message': str(e)}, status=400)


# ---------------------------------------------------------------------------
# PDF Number Canvas (for page numbers)
# ---------------------------------------------------------------------------
class _NumberedCanvas(canvas.Canvas):
    """ReportLab canvas that adds page numbers and community branding to every page."""

    def __init__(self, *args, community_name='', **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []
        self.community_name = community_name

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        total_pages = len(self._saved_page_states)
        for page_num, state in enumerate(self._saved_page_states, start=1):
            self.__dict__.update(state)
            self._draw_footer(page_num, total_pages)
            canvas.Canvas.showPage(self)
        canvas.Canvas.save(self)

    def _draw_footer(self, page_num, total_pages):
        self.saveState()
        self.setFont('Helvetica', 8)
        self.setFillColor(colors.HexColor('#6B7280'))
        footer_text = f'Generated by {self.community_name} Finance System  ·  Page {page_num}/{total_pages}'
        self.drawCentredString(A4[0] / 2, 18 * mm, footer_text)
        # Thin footer line
        self.setStrokeColor(colors.HexColor('#E5E7EB'))
        self.setLineWidth(0.5)
        self.line(20 * mm, 23 * mm, A4[0] - 20 * mm, 23 * mm)
        self.restoreState()


# ---------------------------------------------------------------------------
# Main PDF View
# ---------------------------------------------------------------------------
class MeetingMinutePDFView(LoginRequiredMixin, View):
    """Generates and streams a styled PDF for a meeting minute."""

    def get(self, request, pk, *args, **kwargs):
        minute = get_object_or_404(
            MeetingMinute.objects.select_related('meeting__community', 'prepared_by'),
            pk=pk
        )
        community = minute.meeting.community

        # Attempt to get the digital signature of the preparer
        signature_b64 = ''
        preparer_role = ''
        try:
            membership = Membership.objects.get(community=community, user=minute.prepared_by)
            signature_b64 = membership.digital_signature or ''
            preparer_role = membership.role
        except Membership.DoesNotExist:
            pass

        # Build PDF in memory
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=A4,
            rightMargin=20 * mm,
            leftMargin=20 * mm,
            topMargin=25 * mm,
            bottomMargin=30 * mm,
        )

        # Colour palette
        C_DARK    = colors.HexColor('#1E293B')
        C_PRIMARY = colors.HexColor('#0284C7')
        C_ACCENT  = colors.HexColor('#10B981')
        C_MUTED   = colors.HexColor('#6B7280')
        C_ROW_ALT = colors.HexColor('#F8FAFC')
        C_BORDER  = colors.HexColor('#E2E8F0')
        C_WHITE   = colors.white
        C_GOLD    = colors.HexColor('#F59E0B')

        styles = getSampleStyleSheet()
        W = A4[0] - 40 * mm  # usable width

        def _style(name, parent='Normal', **kw):
            s = ParagraphStyle(name, parent=styles[parent], **kw)
            return s

        # Custom styles
        st_community = _style('CommunityName', fontSize=18, fontName='Helvetica-Bold',
                              textColor=C_DARK, alignment=TA_CENTER, spaceAfter=2)
        st_tagline = _style('Tagline', fontSize=9, fontName='Helvetica',
                            textColor=C_PRIMARY, alignment=TA_CENTER, spaceAfter=4)
        st_title = _style('DocTitle', fontSize=13, fontName='Helvetica-Bold',
                          textColor=C_DARK, alignment=TA_CENTER, spaceAfter=10,
                          spaceBefore=4)
        st_section = _style('SectionHead', fontSize=11, fontName='Helvetica-Bold',
                            textColor=C_DARK, spaceBefore=10, spaceAfter=4)
        st_body = _style('Body', fontSize=9, fontName='Helvetica',
                         textColor=C_DARK, leading=14, spaceAfter=4,
                         alignment=TA_JUSTIFY)
        st_label = _style('TableLabel', fontSize=9, fontName='Helvetica-Bold',
                          textColor=C_PRIMARY)
        st_value = _style('TableValue', fontSize=9, fontName='Helvetica',
                          textColor=C_DARK)
        st_sig_name = _style('SigName', fontSize=10, fontName='Helvetica-Bold',
                             textColor=C_DARK, spaceBefore=4)
        st_sig_meta = _style('SigMeta', fontSize=8, fontName='Helvetica',
                             textColor=C_MUTED, leading=12)

        story = []

        # ── HEADER ──────────────────────────────────────────────────────────
        logo_cell_left  = ''
        logo_cell_right = ''
        try:
            if community.logo:
                # Eagerly load image via PIL into BytesIO so any IO error is
                # caught here (inside try/except), not lazily inside doc.build().
                pil_logo = PILImage.open(community.logo.path).convert('RGBA')
                buf_left = io.BytesIO()
                pil_logo.save(buf_left, format='PNG')
                buf_left.seek(0)
                buf_right = io.BytesIO()
                pil_logo.save(buf_right, format='PNG')
                buf_right.seek(0)
                # Two separate RLImage instances – sharing one between cells breaks layout
                logo_cell_left  = RLImage(buf_left,  width=18 * mm, height=18 * mm)
                logo_cell_right = RLImage(buf_right, width=18 * mm, height=18 * mm)
        except Exception:
            pass  # Logo missing/corrupt – render header without it

        center_header = [
            Paragraph(community.name.upper(), st_community),
            Paragraph('Connect &nbsp;·&nbsp; Inspire &nbsp;·&nbsp; Thrive', st_tagline),
        ]

        header_table = Table(
            [[logo_cell_left, center_header, logo_cell_right]],
            colWidths=[22 * mm, W - 44 * mm, 22 * mm]
        )
        header_table.setStyle(TableStyle([
            ('VALIGN',    (0, 0), (-1, -1), 'MIDDLE'),
            ('ALIGN',     (0, 0), (0, 0),   'LEFT'),
            ('ALIGN',     (1, 0), (1, 0),   'CENTER'),
            ('ALIGN',     (2, 0), (2, 0),   'RIGHT'),
            ('TOPPADDING',    (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        story.append(header_table)
        story.append(HRFlowable(width='100%', thickness=1.5, color=C_PRIMARY, spaceAfter=8))

        # Document title
        story.append(Paragraph('MEETING MINUTES', st_title))
        story.append(HRFlowable(width='100%', thickness=0.5, color=C_BORDER, spaceAfter=10))

        # ── DETAILS TABLE ───────────────────────────────────────────────────
        story.append(Paragraph('Details', st_section))

        # Determine input method
        input_method = 'Voice' if minute.audio_file else 'Text'

        recorded_at = ''
        if minute.created:
            recorded_at = minute.created.strftime('%Y-%m-%d %H:%M:%S')

        preparer_name = ''
        if minute.prepared_by:
            preparer_name = minute.prepared_by.get_full_name or minute.prepared_by.email

        detail_rows = [
            [Paragraph('Title:', st_label),        Paragraph(minute.title or '---', st_value)],
            [Paragraph('Meeting Date:', st_label), Paragraph(str(minute.minute_date) if minute.minute_date else '---', st_value)],
            [Paragraph('Input Method:', st_label), Paragraph(input_method, st_value)],
            [Paragraph('Recorded By:', st_label),  Paragraph(preparer_name or '---', st_value)],
            [Paragraph('Recorded At:', st_label),  Paragraph(recorded_at or '---', st_value)],
        ]

        detail_table = Table(detail_rows, colWidths=[45 * mm, W - 45 * mm])
        detail_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0),  C_ROW_ALT),
            ('BACKGROUND', (0, 2), (-1, 2),  C_ROW_ALT),
            ('BACKGROUND', (0, 4), (-1, 4),  C_ROW_ALT),
            ('TEXTCOLOR',  (0, 0), (0, -1),  C_PRIMARY),
            ('FONTNAME',   (0, 0), (0, -1),  'Helvetica-Bold'),
            ('FONTNAME',   (1, 0), (1, -1),  'Helvetica'),
            ('FONTSIZE',   (0, 0), (-1, -1), 9),
            ('ROWBACKGROUND', (0, 0), (-1, -1), [C_WHITE, C_ROW_ALT]),
            ('GRID',       (0, 0), (-1, -1), 0.5, C_BORDER),
            ('TOPPADDING',    (0, 0), (-1, -1), 6),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ('LEFTPADDING',   (0, 0), (-1, -1), 8),
            ('RIGHTPADDING',  (0, 0), (-1, -1), 8),
        ]))
        story.append(detail_table)
        story.append(Spacer(1, 8))

        # ── SUMMARY ─────────────────────────────────────────────────────────
        if minute.summary:
            story.append(HRFlowable(width='100%', thickness=0.5, color=C_BORDER, spaceAfter=4))
            story.append(Paragraph('Summary', st_section))
            for line in minute.summary.split('\n'):
                line = line.strip()
                if line:
                    story.append(Paragraph(line, st_body))
            story.append(Spacer(1, 6))

        # ── PROFESSIONAL MINUTES ─────────────────────────────────────────────
        if minute.discussions:
            story.append(HRFlowable(width='100%', thickness=0.5, color=C_BORDER, spaceAfter=4))
            story.append(Paragraph('Professional Minutes', st_section))
            for line in minute.discussions.split('\n'):
                line = line.strip()
                if line:
                    story.append(Paragraph(line, st_body))
            story.append(Spacer(1, 6))

        story.append(Spacer(1, 15))
        story.append(Paragraph('Electronic Signature', st_section))
        story.append(HRFlowable(width='100%', thickness=0.5, color=C_BORDER, spaceAfter=8))

        sig_img_cell = ''
        if signature_b64:
            try:
                # Handle data-URI prefix if present
                if ',' in signature_b64:
                    signature_b64 = signature_b64.split(',', 1)[1]
                sig_bytes = base64.b64decode(signature_b64)
                sig_pil = PILImage.open(io.BytesIO(sig_bytes)).convert('RGBA')
                sig_buf = io.BytesIO()
                sig_pil.save(sig_buf, format='PNG')
                sig_buf.seek(0)
                sig_img_cell = RLImage(sig_buf, width=40 * mm, height=25 * mm)
            except Exception:
                sig_img_cell = Paragraph('[ Signature unavailable ]', st_body)
        else:
            sig_img_cell = Paragraph('[ No signature recorded ]', st_body)

        # Right cell: name + meta
        sig_name_text = preparer_name or '---'
        sig_right = [
            Paragraph(sig_name_text, st_sig_name),
            Paragraph(f'Role: {preparer_role}'.lower() if preparer_role else 'Role: ---', st_sig_meta),
            Paragraph(f'Recorded: {recorded_at}', st_sig_meta),
        ]

        sig_table = Table(
            [[
                Table(
                    [['Signed by:'], [sig_img_cell]],
                    colWidths=[55 * mm]
                ),
                sig_right
            ]],
            colWidths=[60 * mm, W - 60 * mm]
        )
        sig_table.setStyle(TableStyle([
            ('BOX',        (0, 0), (-1, -1), 0.5, C_BORDER),
            ('TOPPADDING',    (0, 0), (-1, -1), 8),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
            ('LEFTPADDING',   (0, 0), (-1, -1), 8),
            ('VALIGN',     (0, 0), (-1, -1), 'MIDDLE'),
        ]))
        story.append(sig_table)

        # Build PDF
        def _make_canvas(*args, **kwargs):
            return _NumberedCanvas(*args, community_name=community.name, **kwargs)

        doc.build(story, canvasmaker=_make_canvas)

        # Stream response
        buffer.seek(0)
        safe_title = (minute.title or 'meeting-minutes').replace(' ', '_')[:40]
        filename = f'minutes_{safe_title}.pdf'
        response = HttpResponse(buffer, content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response
